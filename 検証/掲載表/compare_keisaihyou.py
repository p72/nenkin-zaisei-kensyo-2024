#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
財政検証結果レポートの掲載表と通し実行の結果を照合する
======================================================
公表値は厚生労働省「令和6(2024)年財政検証結果レポート」の掲載表
（2024report_back.zip）に**フル精度**で入っている。概要資料の四捨五入値
（56.9% など）ではなく、この生の値と突き合わせるので、桁落ちなしの
完全一致を判定できる。

    # 掲載表を取得して展開
    curl -O https://www.mhlw.go.jp/content/12500000/2024report_back.zip
    unzip -d 掲載表 2024report_back.zip

    # 通し実行（照合したいケースを走らせておく）
    検証/実行/run_pipeline.sh 3001
    検証/実行/run_pipeline.sh 3002
    検証/実行/run_pipeline.sh 3003 1 1 0 2
    検証/実行/run_pipeline.sh 3004 1 1 0 3

    python3 検証/掲載表/compare_keisaihyou.py --keisaihyou 掲載表

照合する表:

| 掲載表 | 内容 | 計算値の出どころ |
|---|---|---|
| 第3-7-29表 | 厚生年金給付費の将来見通し | ⑤ `90nenbe.*_08tou.csv` |
| 第3-7-33表 | 基礎年金給付費の将来見通し | ④ `kekka*a.csv` |
| 第3-7-34表 | 基礎年金拠出金の将来見通し | ④ `kekka*a.csv` |
| 第3-7-35表 | 基礎年金交付金の将来見通し | ④ `kekka*a.csv` |

照合対象は**名目額**の列のみ。各表の第5列（E列）は「賃金上昇率で2024年度
価格に換算した額」という導出値なので、--deflated で別建てに検証する。
"""
import argparse
import collections
import glob
import os
import re
import sys

# --------------------------------------------------------------------------
# 掲載表の定義
# --------------------------------------------------------------------------
# 列は0起点。COL_YEAR は西暦、COL_TOTAL は名目の合計、ITEMS は内訳。
# 内訳の値は「掲載表の列番号のタプル」で、複数なら足して1項目として扱う
# （例: 老齢厚生年金 = 老齢相当 + 通老相当）。
TABLES = {
    '第3-7-29表': {
        'pat': '7-29',
        'title': '厚生年金給付費',
        'source': 'shushi',
        'col_year': 2, 'col_total': 3,
        'items': [('老齢', (5, 6)), ('障害', (7,)), ('遺族', (8,))],
        'sheets': {'高成長実現ケース': '3001', '成長型経済移行・継続ケース': '3002',
                   '過去30年投影ケース': '3003', '１人当たりゼロ成長ケース': '3004'},
    },
    '第3-7-33表': {
        'pat': '7-33',
        'title': '基礎年金給付費',
        'source': 'kekka', 'block': '基礎年金給付費（新法＋旧法）',
        # kekka のブロックは ,合計(制度計,国年,厚年,国共,地共,私学),老齢(...),障害(...),遺族(...)
        'kekka_total': 1, 'kekka_items': [('老齢', (7,)), ('障害', (13,)), ('遺族', (19,))],
        'col_year': 2, 'col_total': 3,
        'items': [('老齢', (5,)), ('障害', (6,)), ('遺族', (7,))],
        'sheets': {'給付費（高成長）': '3001', '給付費（成長移行）': '3002',
                   '給付費（過去投影）': '3003', '給付費（ゼロ成長）': '3004'},
    },
    '第3-7-34表': {
        'pat': '7-34',
        'title': '基礎年金拠出金',
        'source': 'kekka', 'block': '基礎年金拠出金',
        # ,単価,合計,国年,厚年,国共,地共,私学
        'kekka_total': 2,
        'kekka_items': [('国民年金', (3,)), ('被用者年金計', (4, 5, 6, 7)),
                        ('厚生年金', (4,)), ('共済組合', (5, 6, 7))],
        'col_year': 2, 'col_total': 3,
        'items': [('国民年金', (5,)), ('被用者年金計', (6,)),
                  ('厚生年金', (7,)), ('共済組合', (8,))],
        'sheets': {'拠出金（高成長）': '3001', '拠出金（成長移行）': '3002',
                   '拠出金（過去投影）': '3003', '拠出金（ゼロ成長）': '3004'},
    },
    '第3-7-35表': {
        'pat': '7-35',
        'title': '基礎年金交付金',
        'source': 'kekka', 'block': '基礎年金交付金',
        'kekka_total': 1,
        'kekka_items': [('国民年金', (2,)), ('被用者年金計', (3, 4, 5, 6)),
                        ('厚生年金', (3,)), ('共済組合', (4, 5, 6))],
        'col_year': 2, 'col_total': 3,
        'items': [('国民年金', (5,)), ('被用者年金計', (6,)),
                  ('厚生年金', (7,)), ('共済組合', (8,))],
        'sheets': {'交付金（高成長）': '3001', '交付金（成長移行）': '3002',
                   '交付金（過去投影）': '3003', '交付金（ゼロ成長）': '3004'},
    },
}

OKU = 1e4      # 億円 → 兆円
EN = 1e12      # 円   → 兆円


def sniff(path):
    """出力の文字コードを見分ける。原本を UTF-8 化してビルドすれば出力も
    UTF-8、原本のまま EUC-JP でビルドすれば出力も EUC-JP になる（§12.2）。"""
    head = open(path, 'rb').readline()
    for enc in ('utf-8', 'euc_jp'):
        try:
            head.decode(enc)
            return enc
        except UnicodeDecodeError:
            continue
    raise SystemExit(f"文字コードを判別できません: {path}")


# --------------------------------------------------------------------------
# ⑤収支計算の出力（90nenbe）から厚生年金給付費を作る
# --------------------------------------------------------------------------
def read_shushi(path):
    """90nenbe の【スライド調整後】区画を {(西暦, 区分): 兆円} に集める。
    給付費合計 = 1比例 + 2他計。2他計 は 3定額〜7最保 の小計なので、
    毎回その関係を検算してから使う（小計を二重に数えないため）。"""
    enc = sniff(path)
    agg = collections.defaultdict(float)
    parts = collections.defaultdict(float)
    subtotal = collections.defaultdict(float)
    after = seen = False
    name = {'1老齢': '老齢', '2障害': '障害', '3遺族': '遺族'}
    for raw in open(path, 'rb'):
        line = raw.decode(enc, 'replace')
        if line.startswith('【'):
            after, seen = ('調整後' in line), True
            continue
        if not after:
            continue
        f = [x.strip() for x in line.split(',')]
        if len(f) < 8:
            continue
        try:
            nendo, yen = int(f[0]), float(f[7])
        except ValueError:
            continue
        key = (f[1], f[2], f[3], f[4], f[6])
        if f[5] == '2他計':
            subtotal[key] += yen
        elif f[5] != '1比例':
            parts[key] += yen
        if f[5] in ('1比例', '2他計') and f[2] in name:
            agg[(nendo + 2000, name[f[2]])] += yen / OKU

    if not seen:
        raise SystemExit(f"【スライド調整前/後】の区切りがありません: {path}")
    if not agg:
        raise SystemExit(f"【スライド調整後】区画が空です: {path}")
    bad = [k for k in subtotal
           if abs(subtotal[k] - parts[k]) > 1e-6 * max(1.0, abs(subtotal[k]))]
    if bad:
        raise SystemExit(f"2他計が3定額〜7最保の小計になっていない: "
                         f"{len(bad)}件 例 {bad[0]}")
    return agg


# --------------------------------------------------------------------------
# ④基礎年金の出力（kekka*a.csv）からブロックを切り出す
# --------------------------------------------------------------------------
def read_kekka_block(path, block, total_col, items):
    """kekka*a.csv の指定ブロックを {(西暦, 区分): 兆円} に集める。
    ブロックは「カンマを含まない見出し行」で始まり、年度が連続しなくなる
    ところで終わる。単位は円。"""
    enc = sniff(path)
    lines = open(path, 'rb').read().decode(enc, 'replace').splitlines()
    try:
        start = next(i for i, l in enumerate(lines) if l.strip() == block)
    except StopIteration:
        raise SystemExit(f"ブロック「{block}」が見つかりません: {path}")

    agg = {}
    prev = None
    for l in lines[start + 1:]:
        f = [x.strip() for x in l.split(',')]
        if not f or not f[0].isdigit():
            if prev is not None:
                break          # 数値行が始まったあとの非数値行＝ブロックの終わり
            continue           # 見出し行はまだ読み飛ばす
        year = int(f[0])
        if prev is not None and year != prev + 1:
            break
        prev = year
        agg[(year, '合計')] = float(f[total_col]) / EN
        for name, cols in items:
            agg[(year, name)] = sum(float(f[c]) for c in cols) / EN
    if not agg:
        raise SystemExit(f"ブロック「{block}」に数値行がありません: {path}")
    return agg


# --------------------------------------------------------------------------
def find_xlsx(root, pat):
    hits = [p for p in glob.glob(os.path.join(root, '**', '*.xlsx'), recursive=True)
            if pat in os.path.basename(os.path.dirname(p)) or pat in os.path.basename(p)]
    return hits[0] if hits else None


def find_run(spec, case, shushi_dir, bas_dir):
    if spec['source'] == 'shushi':
        hits = sorted(glob.glob(os.path.join(
            shushi_dir, f'90nenbe.{case}-{case}-{case}-{case}-*_08tou.csv')))
    else:
        hits = sorted(glob.glob(os.path.join(
            bas_dir, f'kekka{case}-{case}-{case}-{case}-*a.csv')))
    return hits[0] if hits else None


def read_rv(path):
    """⑤の 03summary の経済前提ブロックから比例改定率を読む。
    このファイルは後続ブロックにも数値行があるので、年度が連続しなくなった
    ところで打ち切る（そうしないと後ろのブロックで上書きされる）。"""
    enc = sniff(path)
    lines = open(path, 'rb').read().decode(enc, 'replace').splitlines()
    hi = next(i for i, l in enumerate(lines) if l.startswith('年度,物価'))
    ci = [c.strip() for c in lines[hi].split(',')].index('比例改定率')
    rv, prev = {}, None
    for l in lines[hi + 1:]:
        f = [x.strip() for x in l.split(',')]
        if not f or not f[0].isdigit():
            break
        y = int(f[0])
        if prev is not None and y != prev + 1:
            break
        prev = y
        rv[y + 2000] = 1.0 + float(f[ci]) / 100.0
    return rv


def check_deflated(args):
    """各掲載表の第5列「賃金上昇率で2024年度価格に換算した額」の換算率を、
    こちらの比例改定率系列と突き合わせる（仕様書 §15.0.2）。

    名目額 D と換算額 E から D/E の年次比を作ると、それが前年度の改定率に
    なっているかを見る。2026年度以降は倍精度の丸めまで一致し、2024年度の
    1要素だけが食い違う。"""
    import openpyxl
    rv = read_rv(find_run({'source': 'shushi'}, '3001', args.shushi_dir, args.bas_dir)
                 .replace('90nenbe.', '03summary.').replace('e_08tou', '_08sum'))

    print("=" * 78)
    print("「2024年度価格」列の換算率（§15.0.2）")
    print("=" * 78)
    print("主張: 換算率の系列はこちらの比例改定率そのもの。")
    print("      ただし2024年度の1要素だけ 1.031 ではなく 1030/999。")
    print()

    ratios = {}
    for tname, spec in TABLES.items():
        xlsx = find_xlsx(args.keisaihyou, spec['pat'])
        if xlsx is None:
            continue
        sheet = next(s for s, c in spec['sheets'].items() if c == '3001')
        wb = openpyxl.load_workbook(xlsx, data_only=True)
        if sheet not in wb.sheetnames:
            continue
        pub = {}
        for row in wb[sheet].iter_rows(values_only=True):
            if len(row) > 4 and isinstance(row[spec['col_year']], (int, float)) \
                    and row[3] and row[4]:
                pub[int(row[spec['col_year']])] = (row[3], row[4])
        if 2025 in pub:
            ratios[tname] = pub[2025][0] / pub[2025][1]
        if tname == '第3-7-29表':
            base = pub

    print("  2025年度の D/E（各表とも同じ値になるはず）")
    for t, r in ratios.items():
        print(f"    {t}  {r:.17f}")
    print(f"    1030/999   = {1030 / 999:.17f}")
    print(f"    こちらの2024年度改定率 = {rv[2024]:.17f}")
    print()

    print(f"  {'年度':<6}{'D/E の年次比':>22}{'こちらの RV[年度−1]':>22}{'相対差':>11}")
    ok = ng = 0
    worst = (0.0, None)
    for y in sorted(base):
        if y - 1 not in base or y - 1 not in rv:
            continue
        r = (base[y][0] / base[y][1]) / (base[y - 1][0] / base[y - 1][1])
        rel = abs(r - rv[y - 1]) / rv[y - 1]
        if y >= 2026:
            if rel > worst[0]:
                worst = (rel, y)
            ok, ng = (ok + 1, ng) if rel < 1e-12 else (ok, ng + 1)
        if y <= 2027 or y == max(base):
            print(f"  {y:<6}{r:>22.16f}{rv[y - 1]:>22.16f}{rel:>11.1e}")
    print()
    print(f"  2026年度以降: 一致 {ok} / 不一致 {ng}   最大相対差 {worst[0]:.2e}"
          f"（{worst[1]}年度）")
    print(f"  2025年度のみ食い違う（相対差 3.0e-05）。その1要素の出所は未特定。")
    print("=" * 78)
    return 0 if ng == 0 and ok > 0 else 1


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--keisaihyou', required=True,
                    help='2024report_back.zip を展開したディレクトリ')
    ap.add_argument('--shushi-dir', default='/suuri/rev2024/emp/rslt/ez_arev/shushi')
    ap.add_argument('--bas-dir', default='/suuri/rev2024/bas/rslt')
    ap.add_argument('--tol', type=float, default=1e-9,
                    help='許容する相対差。既定 1e-9（倍精度の丸め相当）')
    ap.add_argument('--table', help='この掲載表だけ照合（例 第3-7-34表）')
    ap.add_argument('--deflated', action='store_true',
                    help='「2024年度価格」列の換算率を調べる（名目額の照合はしない）')
    args = ap.parse_args()

    if args.deflated:
        return check_deflated(args)

    try:
        import openpyxl
    except ImportError:
        raise SystemExit("openpyxl が必要です: pip install openpyxl")

    total_ok = total_ng = 0
    skipped = []

    for tname, spec in TABLES.items():
        if args.table and args.table != tname:
            continue
        xlsx = find_xlsx(args.keisaihyou, spec['pat'])
        if xlsx is None:
            skipped.append(f"{tname}: 掲載表の xlsx が見つからない（{spec['pat']}）")
            continue
        wb = openpyxl.load_workbook(xlsx, data_only=True)

        print("=" * 78)
        print(f"{tname} {spec['title']}の将来見通し")
        print("=" * 78)

        for sheet, case in spec['sheets'].items():
            if sheet not in wb.sheetnames:
                skipped.append(f"{tname} / {sheet}: 掲載表にシートがない")
                continue
            run = find_run(spec, case, args.shushi_dir, args.bas_dir)
            if run is None:
                skipped.append(f"{tname} / {sheet}: 通し実行の出力がない"
                               f"（試算番号 {case} を走らせていない）")
                continue
            if spec['source'] == 'shushi':
                calc = read_shushi(run)
                items = [('合計', None)] + [(n, None) for n, _ in spec['items']]
            else:
                calc = read_kekka_block(run, spec['block'],
                                        spec['kekka_total'], spec['kekka_items'])
                items = [('合計', None)] + [(n, None) for n, _ in spec['kekka_items']]

            s_ok = s_ng = 0
            worst = (0.0, None)
            for row in wb[sheet].iter_rows(values_only=True):
                if len(row) <= spec['col_year']:
                    continue
                y = row[spec['col_year']]
                if not isinstance(y, (int, float)) or not 2000 <= y <= 2200:
                    continue
                y = int(y)
                pub = {'合計': row[spec['col_total']]}
                for name, cols in spec['items']:
                    if max(cols) >= len(row):
                        continue
                    vals = [row[c] for c in cols]
                    pub[name] = None if all(v is None for v in vals) \
                        else sum(v or 0 for v in vals)
                for name, _ in items:
                    p = pub.get(name)
                    if p is None:
                        continue
                    if spec['source'] == 'shushi' and name == '合計':
                        c = sum(calc.get((y, n), 0.0) for n, _ in spec['items'])
                    else:
                        c = calc.get((y, name))
                    if c is None:
                        continue
                    rel = abs(c - p) / max(abs(p), 1e-12)
                    if rel > worst[0]:
                        worst = (rel, f"{y}年度 {name}")
                    if rel <= args.tol:
                        s_ok += 1
                    else:
                        s_ng += 1
                        if s_ng <= 5:
                            print(f"  不一致 {y}年度 {name:<8} "
                                  f"計算 {c:.10f} / 公表 {p:.10f} 相対差 {rel:.2e}")
            print(f"  {sheet:<26} 一致 {s_ok:>4} / 不一致 {s_ng:>4}   "
                  f"最大相対差 {worst[0]:.1e}"
                  f"{'（' + worst[1] + '）' if worst[1] else ''}")
            total_ok += s_ok
            total_ng += s_ng
        print()

    for s in skipped:
        print(f"（対象外）{s}")
    print()
    print("=" * 78)
    print(f"総合: 一致 {total_ok} 項目 / 不一致 {total_ng} 項目")
    print("=" * 78)
    return 0 if total_ng == 0 and total_ok > 0 else 1


if __name__ == '__main__':
    sys.exit(main())
