#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
財政見通し（詳細結果等の Excel）との照合
========================================
厚生労働省「令和6(2024)年財政検証結果」の**詳細結果等**に入っている財政見通し
の Excel と、通し実行の結果をフル精度で突き合わせます。

    詳細結果等1（通常試算 34本）  https://www.mhlw.go.jp/content/001286770.zip
    詳細結果等2（オプション 92本） https://www.mhlw.go.jp/content/001286771.zip

掲載表（`検証/掲載表/`）は給付費・拠出金・交付金の4表だけですが、こちらは
**収入・支出の全項目と所得代替率・モデル年金額**が四捨五入なしで載っています。

使い方
------
    検証/実行/run_pipeline.sh 3001
    python3 検証/オプション試算/compare_option.py \
        --published '.../01.　人口中位　高成長実現ケース.xlsx'

    TOUGOU=1 YOBI2=001 検証/実行/run_pipeline.sh 3001
    python3 検証/オプション試算/compare_option.py \
        --published '.../21.　調整期間の一致　人口中位　高成長実現ケース.xlsx' \
        --bas-yobi 001

公表側は2種類のレイアウトがあり、シート名で見分けます。

| レイアウト | シート | 出てくる資料 |
|---|---|---|
| 一元化 | 国年年金及び厚年年金 | 調整期間の一致を含むオプション |
| 制度別 | 厚生年金 / 国民年金 | 通常試算とその他のオプション |

こちら側の出どころ
------------------
| 公表の欄 | こちら |
|---|---|
| 厚生年金・一元化の収支 | ⑤ `01shushi.*_08tou.csv`「収支見通し【スライド調整後】」（億円） |
| 国民年金の収支 | ④ `kekka*a.csv`「収支見通し」（円） |
| 所得代替率・モデル年金額 | ⑤ `03summary.*_08sum.csv`「経済前提等」 |
| 国年の一時金・付加年金 | ④ `provide*.csv`「独自給付費等（年度間値）」（円） |

`_08tou` は4制度（厚年・国共済・地共済・私学）の合計です。`_08kou` は第1号
厚生年金被保険者だけなので公表の「厚生年金」には対応しません（2024年度で
59.9兆円 対 68.7兆円）。

調整期間の一致（`Touitu>=1`）では⑤が国民年金の収支まで取り込むので、
`_08tou` が公表の「国民年金及び厚生年金」そのものになります。このとき
国民年金の独自給付は2つの欄に散ります。

    独自給付 = 厚年の独自給付 ＋ 国年の寡婦年金        （shus.c:113）
    事務費   = 厚年の事務費   ＋ 国年の一時金・付加年金・業務勘定への繰入
                                                      （shus.c:382）

公表は「報酬比例」に国年の独自給付を全部入れ、「その他支出」を事務費と
業務勘定への繰入だけにしているので、`provide` から
`一時金納付分＋一時金付加分＋付加年金`（下の X）を取って付け替えます。
"""
import argparse
import glob
import os
import sys

sys.path.insert(0, os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
import suuri_env      # noqa: E402  実行領域の場所（§12.1）

CHO = 1e12            # 円 → 兆円
OKU = 1e-4            # 億円 → 兆円
MAN = 1e-4            # 円 → 万円

UNIFIED_SHEET = '国年年金及び厚年年金'

# 公表シートの列番号。(項目名, 列) で、本体は FIRST_ROW 行から。
UNIFIED_PLAN = {
    'sheet': UNIFIED_SHEET,
    'first_row': 18,
    'rate_row': 11, 'endy_row': 12, 'rate_cols': (14, 15, 16),
    'cols': [('収入合計', 3), ('保険料収入(国年)', 4), ('保険料収入(厚年)', 5),
             ('国庫負担', 6), ('運用収入', 7), ('その他収入', 8),
             ('支出合計', 9), ('基礎年金', 10), ('報酬比例', 11),
             ('その他支出', 12), ('収支差引残', 13), ('年度末積立金', 14),
             ('年度末積立金(2024価格)', 15),
             ('積立度合', 16), ('標準報酬総額', 17),
             ('所得代替率', 19), ('代替率(基礎)', 20), ('代替率(比例)', 21),
             ('手取り収入', 22), ('年金額', 23),
             ('年金額(基礎)', 24), ('年金額(比例)', 25)],
}

BYSYS_COLS_MONEY = [('収入合計', 3), ('保険料収入', 4), ('国庫負担', 5),
                    ('運用収入', 6), ('その他収入', 7), ('支出合計', 8),
                    ('給付費', 9), ('基礎年金拠出金', 10), ('その他支出', 11),
                    ('収支差引残', 12), ('年度末積立金', 13),
                    ('年度末積立金(2024価格)', 14), ('積立度合', 15)]
BYSYS_COLS_RATE = [('所得代替率', 18), ('代替率(基礎)', 19), ('代替率(比例)', 20),
                   ('手取り収入', 21), ('年金額', 22),
                   ('年金額(基礎)', 23), ('年金額(比例)', 24)]

BYSYS_PLAN = {
    'first_row': 17,
    'rate_row': 11, 'endy_row': 12, 'rate_cols': (13, 14, 15),
    'sheets': {
        '厚生年金': BYSYS_COLS_MONEY + [('標準報酬総額', 16)] + BYSYS_COLS_RATE,
        '国民年金': BYSYS_COLS_MONEY + BYSYS_COLS_RATE,
    },
}

PCT = {'所得代替率', '代替率(基礎)', '代替率(比例)'}   # 公表は％、年次欄のみ


# ---------------------------------------------------------------- 公表側

def read_published(path):
    try:
        import openpyxl
    except ImportError:
        raise SystemExit('openpyxl が必要です: pip install openpyxl')
    wb = openpyxl.load_workbook(path, data_only=True)

    if UNIFIED_SHEET in wb.sheetnames:
        layout = '一元化'
        plan = {UNIFIED_PLAN['sheet']: UNIFIED_PLAN['cols']}
        meta = UNIFIED_PLAN
    else:
        layout = '制度別'
        plan = {sn: cols for sn, cols in BYSYS_PLAN['sheets'].items()
                if sn in wb.sheetnames}
        if not plan:
            raise SystemExit(f'知らないレイアウトです: {wb.sheetnames}')
        meta = BYSYS_PLAN

    out = {'path': path, 'layout': layout, 'sheets': {}}
    first = next(iter(plan))
    ws0 = wb[first]
    rc = meta['rate_cols']
    out['代替率'] = {k: ws0.cell(meta['rate_row'], c).value
                     for k, c in zip(('計', '基礎', '比例'), rc)}
    out['終了年度'] = {k: ws0.cell(meta['endy_row'], c).value
                       for k, c in zip(('計', '基礎', '比例'), rc)}
    out['見出し'] = [ws0.cell(r, 2).value for r in (4, 5, 6)]

    for sn, cols in plan.items():
        ws = wb[sn]
        series = {}
        r = meta['first_row']
        while True:
            y = ws.cell(r, 2).value
            if not isinstance(y, int):
                break
            for name, c in cols:
                v = ws.cell(r, c).value
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    v = float(v)
                    if name in PCT:
                        v /= 100.0
                    series.setdefault(name, {})[y] = v
            r += 1
        out['sheets'][sn] = series
    return out


# ---------------------------------------------------------------- こちら側

def sniff(path):
    """出力の文字コードを見分ける（§12.2）。"""
    head = open(path, 'rb').readline()
    for enc in ('utf-8', 'euc_jp'):
        try:
            head.decode(enc)
            return enc
        except UnicodeDecodeError:
            continue
    raise SystemExit(f'文字コードを判別できません: {path}')


def lines_of(path):
    return open(path, 'rb').read().decode(sniff(path), 'replace').splitlines()


def read_block(path, title, header_offset, year_base, lines=None):
    """見出し行の直後に続く年度行を切り出す。戻り値は {年度: {列名: 文字列}}。"""
    L = lines if lines is not None else lines_of(path)
    hits = [i for i, l in enumerate(L) if l.strip() == title]
    if not hits:
        hits = [i for i, l in enumerate(L) if l.startswith(title)]
    if not hits:
        raise SystemExit(f'ブロック「{title}」がありません: {path}')
    start = hits[0]
    cols = [x.strip() for x in L[start + header_offset].split(',')]
    rows = {}
    for l in L[start + header_offset + 1:]:
        f = [x.strip() for x in l.split(',')]
        if not f or not f[0].isdigit():
            if rows:
                break
            continue
        rows[int(f[0]) + year_base] = dict(zip(cols, f))
    if not rows:
        raise SystemExit(f'ブロック「{title}」に年度行がありません: {path}')
    return rows


def find_one(pattern, directory):
    hits = sorted(glob.glob(os.path.join(directory, pattern)))
    if not hits:
        raise SystemExit(
            f'出力が見つかりません: {directory}/{pattern}\n'
            f'先に 検証/実行/run_pipeline.sh を走らせてください。')
    return hits[0]


def read_emp(ver, yobi, shushi_dir):
    """⑤ の 4制度合計（Touitu>=1 なら国年も含む）。単位は兆円。"""
    p = find_one(f'01shushi.{ver}-1120-{yobi}*_08tou.csv', shushi_dir)
    out = {}
    for y, r in read_block(p, '収支見通し【スライド調整後】', 1, 2000).items():
        g = lambda n: float(r[n])                              # noqa: E731
        out[y] = {
            '収入合計':       g('収入計') * OKU,
            '保険料収入':     g('保険料収入') * OKU,
            # 調整期間一致では国年の一時金付加分・付加年金の国庫が
            # 「支援額収入」に入る（shus.c:381）。通常試算では0。
            '国庫負担':       (g('国庫負担') + g('支援額収入')) * OKU,
            '運用収入':       g('運用収入') * OKU,
            'その他収入':     (g('納付金') + g('妻積み分配等')) * OKU,
            '支出合計':       g('支出計') * OKU,
            '独自給付':       g('独自給付') * OKU,
            '基礎年金拠出金': g('基礎年金拠出金') * OKU,
            '事務費':         (g('事務費') + g('支援額支出')) * OKU,
            '収支差引残':     g('収支差') * OKU,
            '年度末積立金':   g('年度末積立金') * OKU,
            '積立度合':       g('積立度合'),
            '標準報酬総額':   g('総報酬総額') * OKU,
        }
    return out


def read_kakaku(ver, bas_yobi, bas_dir):
    """④ kekka の「保険料改定率」列＝2004年度価格への換算率 kakaku[]。

    公表の「2024年度価格」欄は名目額に kakaku[2024] / kakaku[k] を掛けたもの。
    kakaku は 2004年度を1として名目賃金上昇率を積み上げ、2025年度
    （MARUME_NENDO）で小数3桁に丸める（基礎年金/econ.c:312-375）ので、
    kakaku[2024] = 0.999、kakaku[2025] = 1.030 になる。
    掲載表の「2024年度価格」で1年目だけ 1.031 ではなく
    1030/999 = 1.031031031… が使われていたのは、この
    kakaku[2025] / kakaku[2024] そのもの（仕様書 §15.0.2・付録B）。
    """
    p = find_one(f'kekka{ver}-1120-{bas_yobi}a.csv', bas_dir)
    return {y: float(r['保険料改定率'])
            for y, r in read_block(p, '収支見通し', 2, 0).items()}


def read_nat(ver, bas_yobi, bas_dir):
    """④ kekka の国民年金収支。単位は兆円。"""
    p = find_one(f'kekka{ver}-1120-{bas_yobi}a.csv', bas_dir)
    KOKKO = ['国庫（基礎年金）', '国庫（特別国庫）', '国庫（死亡一時金付加分）',
             '国庫（付加年金）', '国庫（旧法寡婦年金免除分）']
    KYUFU = ['死亡一時金納付分', '死亡一時金付加分', '新法寡婦年金',
             '旧法寡婦年金免除分以外', '旧法寡婦年金免除分', '付加年金']
    out = {}
    for y, r in read_block(p, '収支見通し', 2, 0).items():
        g = lambda n: float(r[n])                              # noqa: E731
        shunyu, shishutu = g('収入合計'), g('支出合計')
        out[y] = {
            '収入合計':   shunyu / CHO,
            # 公表は「こども子育て支援特別会計からの繰入」を保険料収入では
            # なくその他収入に入れる（2026年度以降に出てくる）。④の
            # provide が⑤へ渡すときは保険料収入に含めているので、一元化
            # レイアウトの厚年分を作るときは provide 側の合計を使う。
            '保険料収入': (g('保険料収入（国年）')
                           + g('保険料収入（付加年金）')) / CHO,
            '国庫負担':   sum(g(k) for k in KOKKO) / CHO,
            '運用収入':   g('運用収入') / CHO,
            'その他収入': (g('住宅融資債権') + g('妻積み')
                           + g('こども子育て特別会計から繰入')) / CHO,
            'こども繰入': g('こども子育て特別会計から繰入') / CHO,
            '保険料収入(provide)': (g('保険料収入（国年）')
                                    + g('保険料収入（付加年金）')
                                    + g('こども子育て特別会計から繰入')) / CHO,
            '支出合計':   shishutu / CHO,
            '給付費':     sum(g(k) for k in KYUFU) / CHO,
            '基礎年金拠出金': (g('基礎年金拠出金')
                               + g('基礎年金拠出金（特別国庫）')) / CHO,
            'その他支出': g('業務勘定への繰入') / CHO,
            '収支差引残': (shunyu - shishutu) / CHO,
            '年度末積立金': g('年度末積立金') / CHO,
        }
    # 積立度合 = 前年度末積立金 ÷ 当年度支出（§8.1）
    for y in sorted(out):
        if y - 1 in out and out[y]['支出合計']:
            out[y]['積立度合'] = out[y - 1]['年度末積立金'] / out[y]['支出合計']
    return out


def read_rate(ver, yobi, shushi_dir):
    """⑤ 03summary の所得代替率とモデル年金額。代替率は率、年金額は万円。"""
    p = find_one(f'03summary.{ver}-1120-{yobi}*_08sum.csv', shushi_dir)
    L = lines_of(p)
    hi = [i for i, l in enumerate(L) if l.startswith('年度,物価')][0]
    cols = [x.strip() for x in L[hi].split(',')]
    # 「うち比例」「うち基礎」はモデル年金額側にも同名で出るので列番号で取る
    i_rate = cols.index('所得代替率')
    idx = {n: cols.index(n) for n in
           ('モデル年金額（物価割り戻し）', 'うち比例（物価割り戻し）',
            'うち基礎（物価割り戻し）', '可処分所得（物価割り戻し）')}
    out = {}
    for l in L[hi + 1:]:
        f = [x.strip() for x in l.split(',')]
        if not f or not f[0].isdigit():
            if out:
                break
            continue
        y = int(f[0]) + 2000
        out[y] = {
            '所得代替率':   float(f[i_rate]) / 100.0,
            '代替率(比例)': float(f[i_rate + 1]) / 100.0,
            '代替率(基礎)': float(f[i_rate + 2]) / 100.0,
            '年金額':       float(f[idx['モデル年金額（物価割り戻し）']]) * MAN,
            '年金額(比例)': float(f[idx['うち比例（物価割り戻し）']]) * MAN,
            '年金額(基礎)': float(f[idx['うち基礎（物価割り戻し）']]) * MAN,
            '手取り収入':   float(f[idx['可処分所得（物価割り戻し）']]) * MAN,
        }
    return out


def read_provide(ver, yobi, bas_dir):
    """④ provide の「独自給付費等（年度間値）」から X を作る（兆円）。
    X = 一時金納付分 ＋ 一時金付加分 ＋ 付加年金。無ければ全年度0。"""
    hits = sorted(glob.glob(os.path.join(
        bas_dir, '..', 'data', f'provide{ver}-{yobi}.csv')))
    if not hits:
        return {}
    rows = read_block(hits[0], '独自給付費等（年度間値）', 1, 2000)
    return {y: (float(r['一時金納付分']) + float(r['一時金付加分'])
                + float(r['付加年金'])) / CHO
            for y, r in rows.items()}


# ---------------------------------------------------------------- 組み立て

def deflate(nominal, kakaku):
    """名目額を2024年度価格へ。係数は kakaku[2024] / kakaku[k]。"""
    base = kakaku.get(2024)
    if base is None:
        return {}
    return {y: v * base / kakaku[y] for y, v in nominal.items()
            if kakaku.get(y)}


def build(layout, emp, nat, rate, X, kakaku):
    """公表のレイアウトに合わせてこちら側の系列を作る。"""
    if layout == '制度別':
        got = {'厚生年金': {}, '国民年金': {}}
        for label in ('収入合計', '保険料収入', '国庫負担', '運用収入',
                      'その他収入', '支出合計', '収支差引残', '年度末積立金',
                      '積立度合', '基礎年金拠出金'):
            got['厚生年金'][label] = {y: v[label] for y, v in emp.items()}
            got['国民年金'][label] = {y: v[label] for y, v in nat.items()
                                      if label in v}
        got['厚生年金']['給付費'] = {y: v['独自給付'] for y, v in emp.items()}
        got['厚生年金']['その他支出'] = {y: v['事務費'] for y, v in emp.items()}
        got['厚生年金']['標準報酬総額'] = {y: v['標準報酬総額']
                                           for y, v in emp.items()}
        got['国民年金']['給付費'] = {y: v['給付費'] for y, v in nat.items()}
        got['国民年金']['その他支出'] = {y: v['その他支出']
                                         for y, v in nat.items()}
        for sn in got:
            got[sn]['年度末積立金(2024価格)'] = deflate(
                got[sn]['年度末積立金'], kakaku)
        for sn in got:
            for label in ('所得代替率', '代替率(基礎)', '代替率(比例)',
                          '手取り収入', '年金額', '年金額(基礎)', '年金額(比例)'):
                got[sn][label] = {y: v[label] for y, v in rate.items()}
        return got

    # 一元化: ⑤の _08tou が公表表そのもの。国年の独自給付だけ付け替える。
    u = {}
    for label, src in (('収入合計', '収入合計'), ('国庫負担', '国庫負担'),
                       ('運用収入', '運用収入'), ('その他収入', 'その他収入'),
                       ('支出合計', '支出合計'), ('基礎年金', '基礎年金拠出金'),
                       ('収支差引残', '収支差引残'),
                       ('年度末積立金', '年度末積立金'),
                       ('積立度合', '積立度合'),
                       ('標準報酬総額', '標準報酬総額')):
        u[label] = {y: v[src] for y, v in emp.items()}
    # 公表は保険料収入を国年・厚年に分けるが、⑤は合算しか出さない。
    # 国年分は④から取れるので、それを引いて厚年分を作る。
    u['保険料収入(国年)'] = {y: v['保険料収入'] for y, v in nat.items()}
    u['保険料収入(厚年)'] = {y: emp[y]['保険料収入'] - nat[y]['保険料収入(provide)']
                             for y in emp if y in nat}
    u['その他収入'] = {y: v['その他収入'] + nat[y]['こども繰入']
                       for y, v in emp.items() if y in nat}
    u['報酬比例'] = {y: v['独自給付'] + X.get(y, 0.0) for y, v in emp.items()}
    u['その他支出'] = {y: v['事務費'] - X.get(y, 0.0) for y, v in emp.items()}
    u['年度末積立金(2024価格)'] = deflate(u['年度末積立金'], kakaku)
    for label in ('所得代替率', '代替率(基礎)', '代替率(比例)',
                  '手取り収入', '年金額', '年金額(基礎)', '年金額(比例)'):
        u[label] = {y: v[label] for y, v in rate.items()}
    return {UNIFIED_SHEET: u}


# ---------------------------------------------------------------- 照合

def compare(pub, got, rate, tol, atol, show):
    print('=' * 78)
    for h in pub['見出し']:
        if h:
            print(' ', str(h).replace('　', ' '))
    print(f'  レイアウト: {pub["layout"]}')
    print('=' * 78)

    print('\n[所得代替率（給付水準の調整終了後）]')
    fin = rate[max(rate)] if rate else {}
    bad_head = 0
    for key, label in (('計', '所得代替率'), ('基礎', '代替率(基礎)'),
                       ('比例', '代替率(比例)')):
        p, g = pub['代替率'].get(key), fin.get(label)
        if not isinstance(p, (int, float)) or g is None:
            print(f'  {key:<4} 公表 {p}  （対照なし）')
            continue
        d = abs(g - p) / max(abs(p), 1e-30)
        ok = d < 1e-12
        bad_head += 0 if ok else 1
        print(f'  {key:<4} こちら {g * 100:>18.13f}%  公表 {p * 100:>18.13f}%'
              f'   相対差 {d:.1e}  {"一致" if ok else "★不一致"}')
    print('  給付水準の調整終了年度（公表）  '
          + ' / '.join(f'{k}={pub["終了年度"].get(k)}' for k in ('計', '基礎', '比例')))

    total = miss = 0
    worst = (0.0, None)
    left = show
    print()
    for sn, series in got.items():
        pubs = pub['sheets'].get(sn)
        if pubs is None:
            continue
        print(f'[{sn}]')
        print(f'  {"項目":<20}{"年度":>6}{"一致":>6}{"不一致":>7}{"最大相対差":>12}')
        order = ([c[0] for c in UNIFIED_PLAN['cols']] if pub['layout'] == '一元化'
                 else [c[0] for c in BYSYS_PLAN['sheets'][sn]])
        for label in order:
            pv, gv = pubs.get(label, {}), series.get(label, {})
            years = sorted(set(pv) & set(gv))
            if not years:
                print(f'  {label:<20}{"（対照なし）":>25}')
                continue
            n = ng = 0
            mx = (0.0, None)
            for y in years:
                n += 1
                a = abs(gv[y] - pv[y])
                d = a / max(abs(pv[y]), 1e-30)
                if d > mx[0]:
                    mx = (d, y)
                # 収支差引残のように大きな2数の差になる欄は、値そのものが
                # 小さいと相対差が桁落ちで膨らむ。絶対差でも見て、
                # どちらかが許容内なら一致とする（1e-9兆円 = 1000円）。
                if d < tol or a < atol:
                    ng += 1
            total += n
            miss += n - ng
            if mx[0] > worst[0]:
                worst = (mx[0], f'{sn}/{label}/{mx[1]}年度')
            print(f'  {label:<20}{n:>6}{ng:>6}{n - ng:>7}{mx[0]:>12.1e}'
                  f'{"" if ng == n else "  ★"}')
            if ng != n and left > 0:
                for y in years:
                    a = abs(gv[y] - pv[y])
                    d = a / max(abs(pv[y]), 1e-30)
                    if d >= tol and a >= atol:
                        print(f'      {y}  こちら {gv[y]:>20.12f}'
                              f'  公表 {pv[y]:>20.12f}  相対差 {d:.2e}')
                        left -= 1
                        if left <= 0:
                            break
        print()

    print('=' * 78)
    print(f'  照合 {total} 項目、不一致 {miss} 件'
          f'（許容 相対 {tol:.0e} または絶対 {atol:.0e}）')
    if worst[1]:
        print(f'  最大相対差 {worst[0]:.2e}  @ {worst[1]}')
    print('=' * 78)
    return 0 if (miss == 0 and bad_head == 0) else 1


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--published', required=True, help='公表 xlsx のパス')
    ap.add_argument('--case', default='3001', help='試算番号（既定 3001）')
    ap.add_argument('--econ', default=None, help='経済前提番号（既定 試算番号）')
    ap.add_argument('--waku', default=None, help='外枠番号（既定 試算番号）')
    ap.add_argument('--yobi', default='000', help='⑤の予備番号（既定 000）')
    ap.add_argument('--bas-yobi', default=None,
                    help='④の予備番号（既定 --yobi。調整期間一致なら2周目の番号）')
    ap.add_argument('--shushi-dir', default=None)
    ap.add_argument('--bas-dir', default=None)
    ap.add_argument('--tol', type=float, default=1e-12, help='許容相対差')
    ap.add_argument('--atol', type=float, default=1e-9,
                    help='許容絶対差（兆円・％・万円の単位。既定 1e-9）')
    ap.add_argument('--show', type=int, default=8, help='不一致の明細を何件出すか')
    args = ap.parse_args()

    econ = args.econ or args.case
    waku = args.waku or args.case
    bas_yobi = args.bas_yobi or args.yobi
    ver = f'{args.case}-{args.case}-{econ}-{waku}'
    shushi = args.shushi_dir or suuri_env.suuri('emp', 'rslt', 'ez_arev', 'shushi')
    bas = args.bas_dir or suuri_env.suuri('bas', 'rslt')

    pub = read_published(args.published)
    emp = read_emp(ver, args.yobi, shushi)
    nat = read_nat(ver, bas_yobi, bas)
    rate = read_rate(ver, args.yobi, shushi)
    X = read_provide(ver, args.yobi, bas) if pub['layout'] == '一元化' else {}
    kakaku = read_kakaku(ver, bas_yobi, bas)
    got = build(pub['layout'], emp, nat, rate, X, kakaku)
    return compare(pub, got, rate, args.tol, args.atol, args.show)


if __name__ == '__main__':
    sys.exit(main())
