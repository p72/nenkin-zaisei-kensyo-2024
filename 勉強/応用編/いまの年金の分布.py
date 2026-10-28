#!/usr/bin/env python3
"""いま老齢年金をもらう権利がある人の月額の分布（共済を含む）と、国庫負担を廃止して国民全員に一律給付したときの損得（非公式・一部推定）。

分布は2つの公表資料から作る（どちらも令和6年度末＝2025年3月末）。
  ・厚生労働省年金局「令和6年度 厚生年金保険・国民年金事業の概況」（令和7年12月）参考資料3〜5
      会社員の厚生年金（第1号）と国民年金の、年金月額1万円刻みの受給権者数。PDF なので数字を下に書き写した
  ・社会保障審議会年金数理部会「公的年金財政状況報告－令和6(2024)年度－」参考表（Excel）の 2(3)
      国共済・地共済・私学共済の老齢・退職年金の月額の分布（基礎年金を含まない）と、
      第1号の「加入が短い人」の10万円以上の内訳。実行するとダウンロードして .cache/ に置く

人が重ならないように、次の4つを足す。
  加入が長い人（第1号の老齢年金）／加入が短い人（第1号の通算老齢年金・25年未満）／
  共済で加入が長い人（退年相当・老齢相当）／国民年金だけの人（第1号も共済も持たない人）
推定で置いたところ（記事の「推定で置いたところ」と同じ）
  ・報酬比例＝年金月額−基礎年金。基礎年金は全員同じ額 B とし、B は Σmax(月額−B, 0) が公表の厚生年金部分
    （受給権者の年金総額）に合うように決める。共済の月額には基礎年金が入っていないので B を足す
  ・共済で加入が長い人のうち「基礎のみ・共済あり」（第1号を持たない人）を超える分は、会社員の期間もあって
    第1号の「加入が短い人」にも数えられているとみて、その人数を「加入が短い人」から比例で差し引く
  ・共済で加入が短い人（月1〜2万円）は、ほかの年金で数えられているとみて足さない
  ・開いた最後の区分（国民年金だけの7万円以上など）は、公表の平均から逆算した区分に置く

損得は 世代別の損得.py・得する人の割合.py と同じ定義（2027年度からの一律給付と報酬比例のカットを、生涯の
現在価値で比べる。個人単位、遺族年金は入れない）。いまの受給者では「報酬比例が月 thr 万円を超えると損」になる。

使い方（リポジトリの直下で）
  python3 勉強/応用編/いまの年金の分布.py bunpu        # 分布の表だけ（計算プログラムの出力は要らない）
  # 損得も出すときは 得する人の割合.py の準備（通常試算と 701 の 3003・3002）をしておく
  python3 勉強/応用編/いまの年金の分布.py report       # 記事の表
  python3 勉強/応用編/いまの年金の分布.py graph        # 記事の図（勉強/応用編/図/）

openpyxl が要る（pip install openpyxl）。
"""
import hashlib
import importlib.util
import os
import sys
import urllib.request

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
CACHE = os.path.join(ROOT, ".cache", "zaisei_jokyo_r6")
N = 31                                     # 1万円刻みの区分。最後（30）は30万円以上
CASES = (("3003", "過去30年投影"), ("3002", "成長型経済移行・継続"))
REP = 1945                                 # 損得の分かれ目に使う生年度（いまの受給者の真ん中あたり）

# ---- 事業の概況（令和6年度）から書き写した受給権者数（人） ----
# 参考資料3：厚生年金（第1号）の老齢年金。0〜1万円 … 29〜30万円、30万円以上
KM = [30446, 10257, 5404, 5185, 14747, 39134, 134214, 230186, 260278, 260099, 298838, 376357, 456689, 549337, 657775,
      764713, 853718, 926462, 955327, 913998, 820204, 682702, 509842, 341191, 224720, 147563, 92856, 54156, 29810, 14935, 18801]
KF = [12953, 3880, 29993, 63025, 61945, 69313, 180892, 348764, 541901, 751358, 813990, 695128, 522466, 374169, 271489,
      200322, 147604, 105489, 71561, 48617, 33387, 21931, 14116, 8813, 5491, 3233, 1811, 927, 479, 223, 482]
# 参考資料5：厚生年金（第1号）の通算老齢年金・25年未満。0〜1万円 … 9〜10万円、10万円以上
TM = [24039, 62100, 108646, 154369, 219340, 300262, 475303, 621374, 515755, 376070, 527538]
TF = [400511, 353878, 350272, 450653, 799354, 1338108, 2020475, 2426842, 1743529, 789430, 456650]
# 参考資料4・5：国民年金の「基礎のみ共済なし・旧国年」。25年以上（〜7万円以上）と25年未満（〜5万円以上）
NM = [1183, 7154, 28189, 85509, 142939, 151682, 265747, 148474]
NF = [18340, 64351, 171053, 563371, 628734, 760587, 934042, 471390]
SM = [9779, 22708, 16269, 8850, 2247, 283]
SF = [34897, 64306, 46083, 22190, 7677, 1050]
AVG = {"N男": 57309, "N女": 53746, "S男": 20379, "S女": 19704}     # 同じ表の平均年金月額（円）
# 「基礎のみ」と「基礎のみ共済なし」の差＝共済の期間があって第1号を持たない人（参考資料4・5の合計の差）
KIHON_KYO = (5608982 - 4442745) + (238718 - 236339)
# 表10：第1号の受給権者の年金総額（厚生年金部分・年額・億円）
HIREI_TOTAL = {"長": 178071, "短": 26877}
N_L, N_S = 16085696, 14514498              # 参考資料3・5の合計

# ---- 財政状況報告（令和6年度）の参考表 ----
XLSX = {
    "厚年1号": ("001720343.xlsx", "aea9bb5bab425d3392b2905288a2b789112f27ced520f29701a9b87f655ff30e"),
    "国共済": ("001720344.xlsx", "c6bd0711af10df52bb9402a38cb21ff9e62bb225f8a128b22b55bf706f38d649"),
    "地共済": ("001720345.xlsx", "5ac683b51f109370692232baf368fbab298a5c467e325f17c4c5273a223eb25c"),
    "私学共済": ("001720346.xlsx", "2ed333b8d31591fe66ff5450ad2561eb97e835e9cdf4cd5a136f4b6b3e506993"),
}
URL = "https://www.mhlw.go.jp/content/12501000/"


def _get(name):
    """参考表を .cache/ に取ってきて SHA256 を照合する。パスを返す"""
    fn, sha = XLSX[name]
    p = os.path.join(CACHE, fn)
    if not os.path.exists(p):
        os.makedirs(CACHE, exist_ok=True)
        with urllib.request.urlopen(URL + fn, timeout=60) as r, open(p, "wb") as f:
            f.write(r.read())
    h = hashlib.sha256(open(p, "rb").read()).hexdigest()
    if h != sha:
        sys.exit(f"★ {fn} の SHA256 が違う（差し替えられたかもしれない。中身を確かめること）: {h}")
    return p


def _num(x):
    return float(x) if isinstance(x, (int, float)) else 0.0


def _read(name):
    """2(3) 老齢・退職年金受給権者年金月額の分布 → [{"長男","長女","短男","短女": [区分の人数]}]（ブロックごと）。
    人数は 割合（％）× 合計（千人）で出す（千人の列は丸めが粗いため）"""
    import openpyxl
    wb = openpyxl.load_workbook(_get(name), read_only=True, data_only=True)
    out = []
    for sh in [s for s in wb.sheetnames if s.startswith("２(3)")]:
        bins = []
        for r in wb[sh].iter_rows(values_only=True):
            r = [x for x in r if x is not None]
            if not r:
                continue
            s0 = str(r[0])
            if s0[:1] in ("①", "②", "《"):
                bins = []
            elif isinstance(r[0], (int, float)) and len(r) > 1 and r[1] == "～":
                bins.append(r[3:] if r[0] < 30 else r[2:])
            elif s0.startswith("合") and bins:
                if name == "厚年1号":            # 千人・％・'-' の3つ組
                    tot = [_num(r[i]) for i in (1, 4, 10, 13)]
                    pct = lambda v, k: _num(v[3 * k + 1])       # noqa: E731
                else:                            # 千人・％ の2つ組
                    tot = [_num(r[i]) for i in (1, 3, 7, 9)]
                    pct = lambda v, k: _num(v[2 * k + 1])       # noqa: E731
                out.append({key: [pct(v, k) / 100 * t * 1000 for v in bins]
                            for key, k, t in zip(("長男", "長女", "短男", "短女"), (0, 1, 3, 4), tot)})
                bins = []
    return out


def _kokunen(sx):
    """国民年金だけの人（25年以上＋25年未満）。開いた最後の区分は公表の平均から逆算した区分に置く"""
    v = [0.0] * N
    for cnt, avg in (((NM, NF)[sx == "女"], AVG["N" + sx]), ((SM, SF)[sx == "女"], AVG["S" + sx])):
        last = len(cnt) - 1
        rest = avg / 1e4 * sum(cnt) - sum(c * (i + .5) for i, c in enumerate(cnt[:-1]))
        for i, c in enumerate(cnt[:-1]):
            v[i] += c
        v[int(max(rest / cnt[-1], last + .5))] += cnt[-1]
    return v


def _pts(v, n=50):
    """区分の中を一様に n 点ずつ（月額・万円, 人数）。30万円以上は30〜33万円とする"""
    for i, c in enumerate(v):
        if c > 0:
            lo, hi = i, (i + 1 if i < N - 1 else 33)
            for k in range(n):
                yield lo + (hi - lo) * (k + .5) / n, c / n


def _kiso(groups, target):
    """一律の基礎年金 B（万円）。Σmax(T−B,0)/人数 ＝ target になるよう二分法で決める"""
    P = [p for v in groups for p in _pts(v)]
    tot = sum(w for _, w in P)
    lo, hi = 0.0, 10.0
    for _ in range(60):
        b = (lo + hi) / 2
        if sum(w * max(t - b, 0) for t, w in P) / tot > target:
            lo = b
        else:
            hi = b
    return b


_BUNPU = {}


def bunpu():
    """{"P": {性: [(グループ, 年金月額, 報酬比例, 人数)]}, "B": 基礎年金, ...}"""
    if _BUNPU:
        return _BUNPU
    # 第1号・加入が短い人：0〜9万円は概況の実数、10万円以上は財政状況報告の内訳の割合で割る
    k1 = _read("厚年1号")[0]
    S1 = {}
    for sx, base in (("男", TM), ("女", TF)):
        w = k1["短" + sx][10:]
        S1[sx] = list(base[:10]) + [base[10] * x / sum(w) for x in w]
    L1 = {"男": list(KM), "女": list(KF)}
    KYO = {"男": [0.0] * N, "女": [0.0] * N}
    kyo_s = 0.0
    for name in ("国共済", "地共済", "私学共済"):
        for blk in _read(name):
            for sx in ("男", "女"):
                KYO[sx] = [a + b for a, b in zip(KYO[sx], blk["長" + sx])]
                kyo_s += sum(blk["短" + sx])
    n_kyo = sum(KYO["男"]) + sum(KYO["女"])
    dup = max(n_kyo - KIHON_KYO, 0.0)
    sc = 1 - dup / (sum(S1["男"]) + sum(S1["女"]))
    B = _kiso([L1["男"], L1["女"]], HIREI_TOTAL["長"] * 1e8 / N_L / 12 / 1e4)
    BS = _kiso([S1["男"], S1["女"]], HIREI_TOTAL["短"] * 1e8 / N_S / 12 / 1e4)
    P = {}
    for sx in ("男", "女"):
        P[sx] = ([("長", t, max(t - B, 0), w) for t, w in _pts(L1[sx])]
                 + [("短", t, max(t - BS, 0), w * sc) for t, w in _pts(S1[sx])]
                 + [("共済", h + B, h, w) for h, w in _pts(KYO[sx])]
                 + [("国年", t, 0.0, w) for t, w in _pts(_kokunen(sx))])
    _BUNPU.update(P=P, B=B, BS=BS, n_kyo=n_kyo, kyo_s=kyo_s, dup=dup, sc=sc,
                  n_kyo_sx={sx: sum(KYO[sx]) for sx in KYO})
    return _BUNPU


GROUPS = (("長", "会社員の厚生年金（加入が長い人）"), ("短", "会社員の厚生年金（加入が短い人）"),
          ("共済", "公務員・私学の共済（加入が長い人）"), ("国年", "国民年金だけ"))


def hist(items, f=lambda t, h: t):
    """(人数, 年額・兆円) を区分ごとに"""
    v, y = [0.0] * N, [0.0] * N
    for _, t, h, w in items:
        x = f(t, h)
        i = min(max(int(x), 0), N - 1)
        v[i] += w
        y[i] += w * x * 12 / 1e8
    return v, y


def _mean(items, f=lambda t, h: t):
    return sum(w * f(t, h) for _, t, h, w in items) / sum(w for *_, w in items)


# ---- 損得 ----
def _load(name, fname):
    s = importlib.util.spec_from_file_location(name, os.path.join(HERE, fname))
    m = importlib.util.module_from_spec(s)
    s.loader.exec_module(m)
    return m


_SON = {}


def sontoku(case):
    """{"cut", "bi"（2027年度の一律給付・月額・万円）, "thr": {(生年度, 性): 報酬比例の分かれ目（月額・万円）}}"""
    if case in _SON:
        return _SON[case]
    G = _load("tokusuru", "得する人の割合.py").G
    S, SV = G._series(case, "unyo"), G._survival(case)
    R = S["R"]
    thr = {}
    for c in range(1930, 1960):
        for sx, key in (("男", "1"), ("女", "2")):
            s = d = 1.0
            bi = ann = 0.0
            for a in range(65, G.AGE_MAX + 1):
                y = c + a
                if a > 65:
                    s *= SV[key](a, y)
                    d /= 1 + R[y]
                bi += d * s * S["ichi"]["zenin"].get(y, 0.0)
                if y >= G.START:
                    ann += d * s * 12
            thr[c, sx] = bi / 1e4 / (S["cut"] * ann)          # 65歳より前の一律給付は無い世代
    k = G.K.shinkyufu(case, "zenin")[2027]
    _SON[case] = dict(cut=S["cut"], bi=k["getsu_r"] / 1e4, thr=thr)
    return _SON[case]


def _son(case, sx):
    return sontoku(case)["thr"][REP, sx]


# ---- 表 ----
def report_bunpu():
    b = bunpu()
    P = b["P"]
    allp = [p for sx in P for p in P[sx]]
    n = sum(p[3] for p in allp)
    _, y = hist(allp)
    print(f"基礎年金の仮定 B＝{b['B']:.2f}万円（加入が短い人だけで合わせると {b['BS']:.2f}万円）")
    print(f"共済で加入が長い人 {b['n_kyo'] / 1e4:,.0f}万人（男{b['n_kyo_sx']['男'] / 1e4:,.0f}・女{b['n_kyo_sx']['女'] / 1e4:,.0f}）、"
          f"うち第1号とのだぶり（推定）{b['dup'] / 1e4:,.0f}万人 → 加入が短い人を {b['sc']:.3f} 倍。"
          f"足していない共済の加入が短い人 {b['kyo_s'] / 1e4:,.0f}万人\n")
    print("| | 男性 | 女性 | 計 | 平均の月額（男／女） | 年額 |")
    print("|---|---|---|---|---|---|")
    for g, nm in GROUPS:
        it = {sx: [p for p in P[sx] if p[0] == g] for sx in P}
        nn = {sx: sum(p[3] for p in it[sx]) for sx in it}
        yy = sum(hist(it[sx])[1][i] for sx in it for i in range(N))
        print(f"| {nm} | {nn['男'] / 1e4:,.0f}万人 | {nn['女'] / 1e4:,.0f}万人 | {(nn['男'] + nn['女']) / 1e4:,.0f}万人 | "
              f"{_mean(it['男']) * 1e4:,.0f}円／{_mean(it['女']) * 1e4:,.0f}円 | {yy:.1f}兆円 |")
    print(f"| 合計 | {sum(p[3] for p in P['男']) / 1e4:,.0f}万人 | {sum(p[3] for p in P['女']) / 1e4:,.0f}万人 | {n / 1e4:,.0f}万人 | "
          f"{_mean(P['男']) * 1e4:,.0f}円／{_mean(P['女']) * 1e4:,.0f}円（計 {_mean(allp) * 1e4:,.0f}円） | {sum(y):.1f}兆円 |\n")
    v = {sx: hist(P[sx])[0] for sx in P}
    va = [a + c for a, c in zip(v["男"], v["女"])]
    print("| 月額 | 男性 | 女性 | 計（人数） | 計（お金） |")
    print("|---|---|---|---|---|")
    for lim in (7, 10, 15, 20, 25):
        cm = 100 * sum(v["男"][:lim]) / sum(v["男"])
        cf = 100 * sum(v["女"][:lim]) / sum(v["女"])
        print(f"| {lim}万円未満 | {cm:.0f}% | {cf:.0f}% | {100 * sum(va[:lim]) / sum(va):.0f}% | {100 * sum(y[:lim]) / sum(y):.0f}% |")

    def med(a):
        s, c = sum(a), 0.0
        for i, x in enumerate(a):
            if c + x >= s / 2:
                return i + (s / 2 - c) / x
            c += x
    print(f"\n真ん中（区分の中は一様）：人数 {med(va):.1f}万円（男{med(v['男']):.1f}・女{med(v['女']):.1f}）、お金 {med(y):.1f}万円")
    for sx in P:
        j = max(range(N), key=lambda i: v[sx][i])
        print(f"いちばん多い区分（{sx}）：{j}〜{j + 1}万円 {v[sx][j] / 1e4:.0f}万人")


def report():
    report_bunpu()
    P = bunpu()["P"]
    for case, cname in CASES:
        s = sontoku(case)
        th = [s["thr"][k] for k in s["thr"]]
        print(f"\n## {cname}（報酬比例 −{s['cut'] * 100:.1f}%、2027年度の一律給付 月{s['bi'] * 1e4:,.0f}円）")
        print(f"報酬比例の分かれ目：1930〜1959年度生まれで {min(th):.2f}〜{max(th):.2f}万円、"
              f"{REP}年度生まれ 男{_son(case, '男'):.2f}・女{_son(case, '女'):.2f}万円\n")
        print("| | 損する人 | 得する人 | 損する人の割合 |")
        print("|---|---|---|---|")
        tl = tn = 0.0
        for sx in P:
            loss = sum(w for _, _, h, w in P[sx] if h > _son(case, sx))
            n = sum(w for *_, w in P[sx])
            tl += loss
            tn += n
            print(f"| {'男性' if sx == '男' else '女性'} | {loss / 1e4:,.0f}万人 | {(n - loss) / 1e4:,.0f}万人 | {100 * loss / n:.0f}% |")
        print(f"| 計 | {tl / 1e4:,.0f}万人 | {(tn - tl) / 1e4:,.0f}万人 | {100 * tl / tn:.0f}% |")
        print("\n| グループ | 損する人の割合（男／女） |")
        print("|---|---|")
        for g, nm in GROUPS:
            r = []
            for sx in P:
                it = [p for p in P[sx] if p[0] == g]
                r.append(f"{100 * sum(p[3] for p in it if p[2] > _son(case, sx)) / sum(p[3] for p in it):.0f}%")
            print(f"| {nm} | {r[0]}／{r[1]} |")
        # 基礎年金を満額（6.78万円）とおいたときとの比べ
        tl2 = 0.0
        for sx in P:
            tl2 += sum(w for g, t, h, w in P[sx]
                       if (max(t - 6.78, 0) if g in ("長", "短") else h) > _son(case, sx))
        print(f"\n（基礎年金を満額6.78万円とおくと、損する人は {tl2 / 1e4:,.0f}万人・{100 * tl2 / tn:.0f}%）")
        print("\n| | いま | 報酬比例カット後 | ＋一律給付 | いちばん多い区分：いま → ＋一律給付 |")
        print("|---|---|---|---|---|")
        for sx in P:
            f1 = lambda t, h: t - s["cut"] * h                      # noqa: E731
            f2 = lambda t, h: t - s["cut"] * h + s["bi"]            # noqa: E731
            v0, v2 = hist(P[sx])[0], hist(P[sx], f2)[0]
            lo = 7 if sx == "男" else 0                             # 男性は厚生年金の山（7万円以上）を見る
            j0 = max(range(lo, N), key=lambda i: v0[i])
            j2 = max(range(lo, N), key=lambda i: v2[i])
            print(f"| {'男性' if sx == '男' else '女性'}の平均 | {_mean(P[sx]) * 1e4:,.0f}円 | {_mean(P[sx], f1) * 1e4:,.0f}円 | "
                  f"{_mean(P[sx], f2) * 1e4:,.0f}円 | {j0}〜{j0 + 1}万円 → {j2}〜{j2 + 1}万円 |")
        n = sum(w for sx in P for *_, w in P[sx])
        cut = sum(w * s["cut"] * h for sx in P for _, _, h, w in P[sx]) * 12 / 1e8
        print(f"\nこの人たちの報酬比例のカット 年{cut:.1f}兆円、受け取る一律給付 年{n * s['bi'] * 12 / 1e8:.1f}兆円")


# ---- 図 ----
COL = {("長", "男"): "#1f4e8c", ("短", "男"): "#5b8fd0", ("共済", "男"): "#2e8b57", ("国年", "男"): "#b7cfee",
       ("長", "女"): "#a83a37", ("短", "女"): "#dd6f6b", ("共済", "女"): "#d4a017", ("国年", "女"): "#f4c0be"}
NOTE_SRC = ("出典：厚生労働省年金局「令和6年度厚生年金保険・国民年金事業の概況」参考資料3〜5、社会保障審議会年金数理部会"
            "「公的年金財政状況報告－令和6(2024)年度－」参考表（厚年第1号・国共済・地共済・私学共済の2(3)）。令和6年度末＝2025年3月末。")


def _note_suitei(b):
    return (f"推定：共済の月額は基礎年金を含まないので{b['B']:.2f}万円を足した（第1号の公表の厚生年金部分に合うように決めた一律の額）。"
            f"共済で加入が長い人{b['n_kyo'] / 1e4:,.0f}万人のうち約{b['dup'] / 1e4:,.0f}万人は第1号の「加入が短い人」にも数えられているとみて、"
            f"その分を差し引いた。共済で加入が短い人（{b['kyo_s'] / 1e4:,.0f}万人・月1〜2万円）は足していない。障害・遺族年金は含まない。")


def graph():
    G = _load("tokusuru", "得する人の割合.py").G
    matplotlib, plt = G._fig_base()
    from matplotlib.patches import Patch
    b = bunpu()
    P = b["P"]
    w = 0.4
    xs = [i + .5 for i in range(N)]

    def axis(ax, last=True):
        ax.set_xlim(0, 31.3)
        ax.set_xticks(range(0, 32, 2))
        ax.grid(axis="y", color="#ddd", lw=0.6)
        ax.spines[["top", "right"]].set_visible(False)
        if last:
            ax.set_xlabel("年金月額（万円、基礎年金を含む・2024年度の物価）　30〜31 は 30万円以上")

    # 図1：人数とお金
    fig, axs = plt.subplots(2, 1, figsize=(14, 11.5), sharex=True)
    allp = [p for sx in P for p in P[sx]]
    n_all = sum(p[3] for p in allp)
    y_all = sum(hist(allp)[1])
    for k, ax in enumerate(axs):
        for sx, off in (("男", -w / 2), ("女", w / 2)):
            bot = [0.0] * N
            for g, _ in GROUPS:
                v = [x / (1e4 if k == 0 else 1) for x in hist([p for p in P[sx] if p[0] == g])[k]]
                ax.bar([x + off for x in xs], v, w, bottom=bot, color=COL[g, sx])
                bot = [a + c for a, c in zip(bot, v)]
        axis(ax, k == 1)
    axs[0].set_ylabel("受給権者数（万人）")
    axs[1].set_ylabel("年金額（兆円/年）")
    axs[0].set_title(f"人数で見る：{n_all / 1e4:,.0f}万人（平均 月{_mean(allp) * 1e4:,.0f}円）", loc="left", fontsize=12)
    axs[1].set_title(f"お金で見る：年{y_all:.1f}兆円（人数×月額×12）", loc="left", fontsize=12)
    axs[0].legend(handles=[Patch(color=COL[g, sx], label=f"{'男性' if sx == '男' else '女性'}・{nm}"
                                 f"（{sum(p[3] for p in P[sx] if p[0] == g) / 1e4:,.0f}万人）")
                           for sx in ("男", "女") for g, nm in GROUPS], ncol=2, fontsize=9.5, loc="upper right")
    fig.suptitle("いま老齢年金をもらう権利がある人の月額の分布：人数とお金（共済を含む・非公式の集計・一部推定）", fontsize=13.5)
    G._note(fig, NOTE_SRC + "\n" + _note_suitei(b))
    fig.tight_layout(rect=(0, 0.07, 1, 0.97))
    out = os.path.join(HERE, "図", "いまの年金の分布.png")
    fig.savefig(out, dpi=150, facecolor="white")
    plt.close(fig)
    print("書き出し:", out)

    # 図2：損する人・得する人
    fig, axs = plt.subplots(2, 1, figsize=(14, 11.5), sharex=True)
    for j, (ax, (case, cname)) in enumerate(zip(axs, CASES)):
        s = sontoku(case)
        txt = []
        tl = tn = 0.0
        for sx, off, cm in (("男", -w / 2, "#2f5f9e"), ("女", w / 2, "#c0504d")):
            th = _son(case, sx)
            gain = hist([p for p in P[sx] if p[2] <= th])[0]
            loss = hist([p for p in P[sx] if p[2] > th])[0]
            ax.bar([x + off for x in xs], [x / 1e4 for x in gain], w, color=cm, alpha=.35)
            ax.bar([x + off for x in xs], [x / 1e4 for x in loss], w, bottom=[x / 1e4 for x in gain], color=cm)
            n = sum(gain) + sum(loss)
            tl += sum(loss)
            tn += n
            txt.append(f"{'男性' if sx == '男' else '女性'}：損 {sum(loss) / 1e4:,.0f}万人（{100 * sum(loss) / n:.0f}%）・得 {sum(gain) / 1e4:,.0f}万人")
        ax.set_title(f"{cname}（報酬比例 −{s['cut'] * 100:.1f}%、一律給付 月{s['bi'] * 1e4:,.0f}円）：損する人 {tl / 1e4:,.0f}万人（{100 * tl / tn:.0f}%）",
                     loc="left", fontsize=12)
        ax.text(0.99, 0.95, "\n".join(txt) + f"\n分かれ目：報酬比例が月{_son(case, '男'):.1f}万円を超えると損",
                transform=ax.transAxes, ha="right", va="top", fontsize=10.5, bbox=dict(boxstyle="round", fc="white", ec="#999"))
        ax.legend(handles=[Patch(color="#2f5f9e", alpha=.35, label="男性・得"), Patch(color="#2f5f9e", label="男性・損"),
                           Patch(color="#c0504d", alpha=.35, label="女性・得"), Patch(color="#c0504d", label="女性・損")],
                  loc="upper left", ncol=2, fontsize=10)
        ax.set_ylabel("受給権者数（万人）")
        ax.set_ylim(0, 320)
        axis(ax, j == 1)
    fig.suptitle("厚生年金の国庫負担を廃止して国民全員に一律給付（財政均衡）：いまの受給者で損する人・得する人（非公式の推定）", fontsize=13.5)
    G._note(fig, "非公式：厚生労働省の試算ではない。損得は2027年度からの一律給付と報酬比例のカットの、生涯の現在価値の差"
            f"（2024年財政検証の計算プログラムに独自のレバーを加えた試算から推定。{REP}年度生まれの分かれ目を使う。個人単位、遺族年金は入れていない）。"
            f"報酬比例＝年金月額−基礎年金（{b['B']:.2f}万円で一律とおいた推定）。国民年金だけの人は全員得。\n" + NOTE_SRC + "\n" + _note_suitei(b))
    fig.tight_layout(rect=(0, 0.085, 1, 0.97))
    out = os.path.join(HERE, "図", "いまの受給者の損得.png")
    fig.savefig(out, dpi=150, facecolor="white")
    plt.close(fig)
    print("書き出し:", out)

    # 図3：報酬比例の一律カットで山がどう動くか（過去30年投影）
    case, cname = CASES[0]
    s = sontoku(case)
    fig, axs = plt.subplots(2, 1, figsize=(14, 10.5), sharex=True)
    for j, (ax, sx, cm) in enumerate(zip(axs, ("男", "女"), ("#2f5f9e", "#c0504d"))):
        for lab, f, col, ls, lw in (("いま", lambda t, h: t, "#888888", "-", 2.6),
                                    (f"報酬比例を −{s['cut'] * 100:.1f}%", lambda t, h: t - s["cut"] * h, cm, "-", 2.2),
                                    (f"＋一律給付 月{s['bi'] * 1e4:,.0f}円", lambda t, h: t - s["cut"] * h + s["bi"], "#222222", "--", 1.8)):
            v = hist(P[sx], f)[0]
            ax.step(list(range(N + 1)), [x / 1e4 for x in v] + [v[-1] / 1e4], where="post", color=col, ls=ls, lw=lw,
                    label=f"{lab}（平均 月{_mean(P[sx], f) * 1e4:,.0f}円）")
        ax.set_title(f"{'男性' if sx == '男' else '女性'}（{sum(p[3] for p in P[sx]) / 1e4:,.0f}万人）", loc="left", fontsize=12)
        ax.set_ylabel("受給権者数（万人）")
        ax.legend(fontsize=10.5, loc="upper right")
        axis(ax, j == 1)
    fig.suptitle(f"報酬比例の一律カット（{cname} −{s['cut'] * 100:.1f}%）で、いまの受給者の月額の山がどう動くか（非公式の推定）", fontsize=13.5)
    G._note(fig, "非公式：厚生労働省の試算ではない。いまの受給者の年金月額に、2027年度からの報酬比例の一律カットと一律給付"
            "（2027年度の額・2024年度の物価）をそのまま当てはめた。基礎年金は変わらない。"
            f"報酬比例＝年金月額−基礎年金（{b['B']:.2f}万円で一律とおいた推定）。共済は月額をまるごと報酬比例とした（旧3階を含む）。\n"
            + NOTE_SRC + "\n" + _note_suitei(b))
    fig.tight_layout(rect=(0, 0.085, 1, 0.97))
    out = os.path.join(HERE, "図", "いまの受給者_報酬比例カット.png")
    fig.savefig(out, dpi=150, facecolor="white")
    plt.close(fig)
    print("書き出し:", out)


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else "report"
    {"bunpu": report_bunpu, "report": report, "graph": graph}[cmd]()
